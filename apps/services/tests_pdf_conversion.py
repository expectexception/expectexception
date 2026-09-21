"""Round-trip fidelity/editability tests for the PDF <-> DOCX conversion
pipeline (apps/services/pdf_utils.py).

A real python-docx document is built, run through the actual soffice /
pdf2docx / PyMuPDF pipeline, and the resulting DOCX is inspected for
structure (heading levels, bold/italic runs, tables) rather than just "did
it not crash". That distinction matters here: the native rebuild's heading
detection used to use fixed size-ratio cutoffs (>=1.6x body for H1, >=1.2x
for H2) that were never checked against a real Word-generated document —
Word's own Heading 1/Heading 2 defaults land at roughly 1.27x/1.18x body
size, so real headings were silently misclassified or dropped to plain
paragraphs. These tests catch that class of regression.

Skipped entirely when soffice/libreoffice isn't on PATH, since the DOCX->PDF
leg (and therefore everything downstream) depends on it. CI's default `test`
job installs Python deps only (see .github/workflows/ci.yml), not
LibreOffice, so this suite is a no-op there and runs for real in any
environment — like this repo's Docker image — that has it installed.
"""

import shutil
import tempfile
import unittest
from pathlib import Path

from django.test import TestCase
from docx import Document

from apps.services.pdf_utils import (
    convert_document_to_pdf,
    convert_pdf_to_docx_native,
    convert_pdf_with_pdf2docx,
)

SOFFICE_AVAILABLE = bool(shutil.which("soffice") or shutil.which("libreoffice"))


def _build_sample_docx(path: str) -> None:
    doc = Document()
    doc.add_heading("Quarterly Report", level=1)
    doc.add_heading("Summary", level=2)
    p = doc.add_paragraph("Revenue grew ")
    p.add_run("18%").bold = True
    p.add_run(" year over year, driven mostly by the ")
    p.add_run("EMEA").italic = True
    p.add_run(" region.")

    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "Region"
    table.cell(0, 1).text = "Revenue"
    table.cell(1, 0).text = "EMEA"
    table.cell(1, 1).text = "$1.1M"

    doc.add_paragraph("End of report.")
    doc.save(path)


@unittest.skipUnless(SOFFICE_AVAILABLE, "soffice/libreoffice not on PATH")
class PdfConversionFidelityTests(TestCase):
    """DOCX -> PDF -> DOCX round trip: structure must survive."""

    def setUp(self):
        self.tmp_dir = tempfile.mkdtemp(prefix="pdf_conv_test_")
        self.addCleanup(shutil.rmtree, self.tmp_dir, ignore_errors=True)
        self.src_docx = str(Path(self.tmp_dir) / "source.docx")
        self.pdf_path = str(Path(self.tmp_dir) / "source.pdf")
        _build_sample_docx(self.src_docx)
        convert_document_to_pdf(self.src_docx, self.pdf_path)

    def test_docx_to_pdf_produces_nonempty_file(self):
        self.assertTrue(Path(self.pdf_path).exists())
        self.assertGreater(Path(self.pdf_path).stat().st_size, 0)

    def test_native_rebuild_preserves_heading_levels(self):
        out_path = str(Path(self.tmp_dir) / "native.docx")
        convert_pdf_to_docx_native(self.pdf_path, out_path)
        doc = Document(out_path)

        headings = {p.text.strip(): p.style.name for p in doc.paragraphs if p.text.strip()}
        self.assertEqual(headings.get("Quarterly Report"), "Heading 1")
        self.assertEqual(headings.get("Summary"), "Heading 2")

    def test_native_rebuild_preserves_bold_and_italic_runs(self):
        out_path = str(Path(self.tmp_dir) / "native.docx")
        convert_pdf_to_docx_native(self.pdf_path, out_path)
        doc = Document(out_path)

        runs = [r for p in doc.paragraphs for r in p.runs]
        self.assertTrue(any(r.text.strip() == "18%" and r.bold for r in runs))
        self.assertTrue(any(r.text.strip() == "EMEA" and r.italic for r in runs))

    def test_native_rebuild_preserves_table(self):
        out_path = str(Path(self.tmp_dir) / "native.docx")
        convert_pdf_to_docx_native(self.pdf_path, out_path)
        doc = Document(out_path)

        self.assertEqual(len(doc.tables), 1)
        cells = [[c.text for c in row.cells] for row in doc.tables[0].rows]
        self.assertEqual(cells, [["Region", "Revenue"], ["EMEA", "$1.1M"]])

    def test_pdf2docx_preserves_table(self):
        out_path = str(Path(self.tmp_dir) / "pdf2docx.docx")
        convert_pdf_with_pdf2docx(self.pdf_path, out_path)
        doc = Document(out_path)

        self.assertEqual(len(doc.tables), 1)
        cells = [[c.text for c in row.cells] for row in doc.tables[0].rows]
        self.assertEqual(cells, [["Region", "Revenue"], ["EMEA", "$1.1M"]])
